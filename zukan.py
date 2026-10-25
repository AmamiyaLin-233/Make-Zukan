import os
from pptx import Presentation
import tkinter
from tkinter import filedialog


class ZukanApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Make Zukan")
        self.root.geometry("445x250")
        self.create_widgets()

    def create_widgets(self):
        # 按钮区域
        frame = tkinter.Frame(self.root)
        frame.grid(row=0, column=0, padx=0, pady=0, sticky="w")

        button1 = tkinter.Button(frame, text=" Start ", command=self.btn1_click)
        button1.grid(row=0, column=1, padx=10, pady=5)

        button2 = tkinter.Button(frame, text=" Open pptx ", command=self.btn2_click)
        button2.grid(row=0, column=2, padx=0, pady=10)

        button3 = tkinter.Button(frame, text=" XXX ", command=self.btn3_click)
        # button3.grid(row=0, column=3, padx=10, pady=10)

        # 路径显示标签
        self.label1 = tkinter.Label(self.root, text="你的图片文件路径")
        self.label1.grid(row=1, column=0, padx=10, pady=0, sticky="w")

        # 文本日志框
        self.textbox = tkinter.Text(self.root, width=60, height=10)
        self.textbox.insert(
            tkinter.END,
            "请选择文件夹后生成。文件名应当以逗号分隔[科属]和[种]。\n",
        )
        self.textbox.grid(row=2, column=0, padx=10, pady=0)

        # 版权信息
        label2 = tkinter.Label(self.root, text="Copyright © 2026 AmamiyaLin All Rights Reserved")
        label2.config(fg="gray")
        label2.grid(row=3, column=0, padx=10, pady=10, sticky="w")

    def generate_catalog_with_placeholders(self, template_path, image_folder):
        # 检查模板文件是否存在
        if not os.path.exists(template_path):
            self.textbox.insert(tkinter.END, f"错误：找不到模板文件 '{template_path}'\n")
            return

        prs = Presentation(template_path)
        slide_layout = prs.slide_layouts[0]
        valid_extensions = (".png", ".jpg", ".jpeg", ".bmp", ".webp")

        # 获取有效图片列表
        image_files = [f for f in os.listdir(image_folder) if f.lower().endswith(valid_extensions)]

        if not image_files:
            self.textbox.insert(tkinter.END, "提示：所选文件夹中没有找到支持的图片文件。\n")
            return

        for img_file in image_files:
            filename = os.path.splitext(img_file)[0]
            s = "".join(ch for ch in filename if not ch.isdigit())
            shuzh, _, hua = s.partition("，")
            if hua == "":
                hua, shuzh = shuzh, hua

            img_path = os.path.join(image_folder, img_file)
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = hua

            for shape in slide.placeholders:
                if shape.placeholder_format.type == 18:  # 图片占位符
                    shape.insert_picture(img_path)
                elif shape.has_text_frame and shape != slide.shapes.title:
                    shape.text = shuzh

        # 使用 os.path.join 安全拼接输出路径
        output_path = os.path.join(image_folder, "output_zukan.pptx")
        prs.save(output_path)
        self.textbox.insert(tkinter.END, f"图鉴生成完毕！已保存至: {output_path}\n")

    def btn1_click(self):
        folder = filedialog.askdirectory(title="选择存放图片的文件夹")
        if folder:  # 用户选择了文件夹（没有点击取消）
            self.label1.config(text=folder)
            template = "template.pptx"
            self.generate_catalog_with_placeholders(template, folder)

    def btn2_click(self):
        folder = self.label1.cget("text")
        output_path = os.path.join(folder, "output_zukan.pptx")

        if os.path.exists(output_path):
            try:
                os.startfile(output_path)  # 注意：os.startfile 仅适用于 Windows
            except Exception as e:
                self.textbox.insert(tkinter.END, f"打开文件失败: {e}\n")
        else:
            self.textbox.insert(tkinter.END, "提示：尚未生成 pptx 文件或选择有效路径。\n")

    def btn3_click(self):
        pass


if __name__ == "__main__":
    main_window = tkinter.Tk()
    app = ZukanApp(main_window)
    main_window.mainloop()
